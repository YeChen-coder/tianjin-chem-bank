from copy import deepcopy
import hashlib
import json
import math
from pathlib import Path
import uuid

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from lxml import etree as E
from PIL import Image

from question_content import compile_formula, validate_bundle
from question_content.storage import ContentError

M='http://schemas.openxmlformats.org/officeDocument/2006/math'
ALIGN={'left':WD_ALIGN_PARAGRAPH.LEFT,'center':WD_ALIGN_PARAGRAPH.CENTER,'right':WD_ALIGN_PARAGRAPH.RIGHT}


def export_paper(bundle, output_path, options=None, asset_resolver=None):
    """Returns output_path, report_path and report; never mutates bundle or input assets.

    asset_resolver is a trusted callable(asset_id)->local file. It may expose
    resolve_source(source_ref) and resolve_preview(asset_id) capabilities.
    """
    return _Exporter(bundle, options or {}, asset_resolver).export(Path(output_path))


class _Exporter:
    def __init__(self,bundle,options,resolver):
        validation=validate_bundle(bundle)
        if not validation['ok']: raise ContentError('VALIDATION_FAILED','导出数据验证失败',details=validation)
        allowed={'include_answers','include_explanations','table_width_policy','fail_on_issues','repeat_table_header','auto_number','answers_position'}
        if not isinstance(options,dict) or options.keys()-allowed: raise ContentError('BAD_REQUEST','导出选项不合法',400)
        for k in ('include_answers','include_explanations','fail_on_issues','repeat_table_header','auto_number'):
            if k in options and type(options[k]) is not bool: raise ContentError('BAD_REQUEST',k+' 必须为布尔值',400)
        if options.get('table_width_policy','scale') not in ('scale','error'): raise ContentError('BAD_REQUEST','table_width_policy 必须为 scale/error',400)
        if options.get('answers_position','end') not in ('end','inline'):raise ContentError('BAD_REQUEST','答案位置无效',400)
        self.bundle=deepcopy(bundle);self.options=options;self.resolver=resolver
        self.assets={a['id']:a for a in bundle['assets']}
        self.report={'schema_version':1,'paper_id':bundle['paper']['id'],'revisions':{q['id']:q['revision'] for q in bundle['questions']},'formula_count':0,'table_count':0,'image_count':0,'issues':deepcopy(bundle['issues']),'native_source_count':0,'image_assets':[],'visual_verified':False}
        known={i['id'] for i in self.report['issues']}
        for q in bundle['questions']:
            for issue in q['issues']:
                if issue['id'] not in known:self.report['issues'].append(deepcopy(issue));known.add(issue['id'])
        self.qid=None
        self.doc=Document();section=self.doc.sections[0]
        section.page_width=Mm(210);section.page_height=Mm(297)
        section.top_margin=section.bottom_margin=Mm(18)
        section.left_margin=section.right_margin=Mm(20)
        self.width=170; self.height=261
        normal=self.doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(11)
        normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
        normal.paragraph_format.space_after=Pt(5)
        normal.paragraph_format.line_spacing=1.25

    def issue(self,code,message,node=None,severity='review'):
        self.report['issues'].append({'id':'export-'+uuid.uuid4().hex,'code':code,'severity':severity,'question_id':self.qid,'node_id':node.get('id') if node else None,'message':message})

    def paragraph(self,parent):
        # Reuse the mandatory first empty paragraph in a freshly cleared cell.
        if parent is not self.doc and hasattr(parent,'paragraphs') and len(parent.paragraphs)==1 and len(parent.paragraphs[0]._p)==0:
            return parent.paragraphs[0]
        return parent.add_paragraph()

    def formula(self,p,n):
        record=None
        source=getattr(self.resolver,'resolve_source',None)
        if source and n.get('source_ref'):record=source(n['source_ref'])
        result=None
        if record and record.get('omml') and (n.get('latex') is None or n.get('latex')==record.get('latex')):
            try:
                native=E.fromstring(record['omml'].encode(),E.XMLParser(resolve_entities=False,no_network=True))
                if native.tag=='{'+M+'}oMathPara':native=native.find('{'+M+'}oMath')
                if native is None or native.tag!='{'+M+'}oMath' or not ''.join(native.itertext()).strip(): raise ValueError('原生公式结构为空')
                p._p.append(deepcopy(native));self.report['formula_count']+=1;self.report['native_source_count']+=1
                if n['status']!='confirmed':self.issue('FORMULA_REVIEW','未确认公式保留来源原生结构',n)
                return
            except (ValueError,E.XMLSyntaxError,TypeError) as exc:self.issue('UNSUPPORTED_SOURCE','来源公式不能恢复：'+str(exc),n)
        if n.get('latex') is not None: result=compile_formula(n['latex'],n.get('source_ref'))
        if result and result['ok']:
            p._p.append(E.fromstring(result['omml'].encode(),E.XMLParser(resolve_entities=False,no_network=True)))
            self.report['formula_count']+=1
            if n['status'] in {'review_required','unsupported'}: self.issue('FORMULA_REVIEW','候选公式需要教师核对',n)
            return
        self.issue('FORMULA_COMPILE_FAILED','公式未编译，保留原内容与来源：'+(n.get('source_ref') or n['id']),n,'error')
        p.add_run('[待核对公式：'+(n.get('latex') or (record or {}).get('text') or n.get('source_ref') or n['id'])+']')

    def picture(self,p,n,width):
        try:
            if not self.resolver: raise ValueError('未提供资产解析器')
            path=Path(self.resolver(n['asset_id']))
            sha=hashlib.sha256(path.read_bytes()).hexdigest()
            if sha!=self.assets[n['asset_id']]['sha256']: raise ValueError('原媒体摘要不一致')
            render_path=path
            try:
                with Image.open(path) as image: iw,ih=image.size
            except (OSError,ValueError):
                preview=getattr(self.resolver,'resolve_preview',None)
                if not preview:raise
                render_path=Path(preview(n['asset_id']))
                with Image.open(render_path) as image:iw,ih=image.size
                self.issue('IMAGE_PREVIEW_FALLBACK','原媒体保留，Word 使用登记的预览图',n)
            requested=n['width_mm'];height=requested*ih/iw
            if abs(height-n['height_mm'])/height>0.02:self.issue('IMAGE_ASPECT_ADJUSTED','图片声明尺寸与原比例不符，按原比例导出',n)
            scale=min(1,width/requested,(self.height-20)/height)
            if scale<1:self.issue('IMAGE_SCALED','图片按页面/单元格可用尺寸等比缩小',n,'info')
            p.add_run().add_picture(str(render_path),width=Mm(requested*scale),height=Mm(height*scale))
            self.report['image_count']+=1;self.report['image_assets'].append({'asset_id':n['asset_id'],'original_sha256':sha,'render_sha256':hashlib.sha256(render_path.read_bytes()).hexdigest()})
        except Exception as exc:
            p.add_run('[图片待处理：'+n['asset_id']+']')
            self.issue('ASSET_MISSING','图片未输出：'+str(exc),n,'error')

    def blocks(self,parent,blocks,width,prefix='',keep_last=False):
        for index,b in enumerate(blocks):
            if index==0 and prefix and b['type']!='paragraph':
                p=self.paragraph(parent);p.add_run(prefix);p.paragraph_format.keep_with_next=True
            if b['type']=='paragraph':
                p=self.paragraph(parent);p.alignment=ALIGN[b.get('align','left')]
                p.paragraph_format.widow_control=True
                if (index+1<len(blocks) and blocks[index+1]['type'] in {'image','table'}) or (keep_last and index==len(blocks)-1):p.paragraph_format.keep_with_next=True
                if index==0 and prefix:p.add_run(prefix)
                for n in b['inlines']:
                    if n['type']=='formula':self.formula(p,n)
                    elif n['type']=='image':self.picture(p,n,width)
                    elif n['type']=='blank':
                        run=p.add_run('_'*max(2,math.ceil(min(n['width_mm'],width)/0.8)))
                        fit=OxmlElement('w:fitText');fit.set(qn('w:val'),str(round(min(n['width_mm'],width)/25.4*1440)))
                        fit.set(qn('w:id'),str(int(hashlib.sha256(n['id'].encode()).hexdigest()[:7],16)))
                        run._r.get_or_add_rPr().append(fit)
                        if n['width_mm']>width:self.issue('BLANK_WIDTH','作答空超过可用宽度',n)
                    else:
                        r=p.add_run(n['text']);style=n.get('style',{})
                        r.bold=style.get('bold',False);r.italic=style.get('italic',False)
                        r.font.subscript=style.get('vertical')=='subscript';r.font.superscript=style.get('vertical')=='superscript'
            elif b['type']=='image':
                p=self.paragraph(parent);p.alignment=ALIGN[b.get('align','left')];self.picture(p,b,width)
            elif b['type']=='unresolved':
                p=self.paragraph(parent);p.add_run('[待核对：'+b['reason']+']')
                self.issue('UNSUPPORTED_SOURCE',b['reason']+'；来源 '+b['source_ref'],b)
                self.blocks(parent,b['fallback_blocks'],width)
            else:self.table(parent,b,width)

    def table(self,parent,b,width):
        self.report['table_count']+=1
        columns=b['columns_mm'];factor=min(1,width/sum(columns))
        if factor<1:
            if self.options.get('table_width_policy','scale')=='error':raise ContentError('TABLE_TOO_WIDE','表格超过页面可用宽度')
            self.issue('TABLE_SCALED','表格列宽按可用宽度同比缩小',b,'info')
        columns=[w*factor for w in columns]
        if parent is self.doc:t=parent.add_table(rows=len(b['rows']),cols=len(columns))
        else:t=parent.add_table(rows=len(b['rows']),cols=len(columns))
        t.style='Table Grid';t.autofit=False
        for col,w in zip(t.columns,columns):col.width=Mm(w)
        for r,row in enumerate(b['rows']):
            if r==0 and self.options.get('repeat_table_header',False):
                prop=OxmlElement('w:tblHeader');t.rows[r]._tr.get_or_add_trPr().append(prop)
            for c,w in enumerate(columns):t.cell(r,c).width=Mm(w)
        for r,row in enumerate(b['rows']):
            for cell in sorted(row['cells'],key=lambda c:c['col_index']):
                c=cell['col_index'];target=t.cell(r,c)
                if cell['row_span']>1 or cell['col_span']>1:target=target.merge(t.cell(r+cell['row_span']-1,c+cell['col_span']-1))
                target.width=Mm(sum(columns[c:c+cell['col_span']]))
                margins=OxmlElement('w:tcMar')
                for side in ('top','bottom','left','right'):
                    margin=OxmlElement('w:'+side);margin.set(qn('w:w'),'100');margin.set(qn('w:type'),'dxa');margins.append(margin)
                target._tc.get_or_add_tcPr().append(margins)
                for child in list(target._tc):
                    if child.tag!=qn('w:tcPr'):target._tc.remove(child)
                target.add_paragraph()
                self.blocks(target,cell['blocks'],max(1,sum(columns[c:c+cell['col_span']])-4))

    def export(self,path):
        title=self.doc.add_paragraph();title.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r=title.add_run(self.bundle['paper']['title']);r.bold=True;r.font.size=Pt(16)
        title.paragraph_format.keep_with_next=True
        qmap={q['id']:q for q in self.bundle['questions']}
        def qnumber(qid):
            return str(self.bundle['paper']['question_ids'].index(qid)+1) if self.options.get('auto_number',False) else ''
        def write_answers(qid):
            q=qmap[qid];self.qid=qid;labels={q['id']:''}
            def targets(items):
                for item in items:labels[item['id']]='（'+item['label']+'）' if item['label'].isdigit() else item['label'];targets(item.get('subquestions',[]))
            targets(q['subquestions']);started=False
            for field,label in (('answers','答案'),('explanations','解析')):
                if not self.options.get('include_'+field,True) or not q[field]:continue
                marker=(qnumber(qid)+'. ' if not started and qnumber(qid) else '')+'【'+label+'】';started=True
                for index,a in enumerate(q[field]):
                    target=labels.get(a['target_id'])
                    if target is None:
                        from question_content.storage import nodes
                        blanks=[n['id'] for n in nodes({k:q[k] for k in ('body','options','subquestions')}) if n.get('type')=='blank']
                        target='第'+str(blanks.index(a['target_id'])+1)+'空：'
                    self.blocks(self.doc,a['blocks'],self.width,(marker if index==0 else '')+(target+' ' if target else ''))
        for qid in self.bundle['paper']['question_ids']:
            for annotation in self.bundle['paper'].get('annotations',[]):
                if annotation['before_question_id']==qid:
                    self.qid=None;self.blocks(self.doc,annotation['blocks'],self.width,keep_last=True)
            q=qmap[qid];self.qid=qid
            self.blocks(self.doc,q['body'],self.width,qnumber(qid)+'. ' if qnumber(qid) else '',keep_last=bool(q['options'] or q['subquestions']))
            for option in q['options']:self.blocks(self.doc,option['blocks'],self.width,option['label']+'. ')
            def subs(items):
                for item in items:
                    label='（'+item['label']+'）' if item['label'].isdigit() else item['label']
                    self.blocks(self.doc,item['blocks'],self.width,label+' ')
                    subs(item.get('subquestions',[]))
            subs(q['subquestions'])
            if self.options.get('answers_position','end')=='inline':write_answers(qid)
        for annotation in self.bundle['paper'].get('annotations',[]):
            if annotation['before_question_id'] is None:
                self.qid=None;self.blocks(self.doc,annotation['blocks'],self.width)
        if self.options.get('answers_position','end')=='end' and (self.options.get('include_answers',True) or self.options.get('include_explanations',True)):
            if any(q['answers'] or q['explanations'] for q in qmap.values()):
                self.doc.add_page_break();p=self.doc.add_paragraph('参考答案与解析');p.runs[0].bold=True
                for qid in self.bundle['paper']['question_ids']:
                    write_answers(qid)
        self.report['ok']=not any(i['severity']=='error' for i in self.report['issues'])
        if self.options.get('fail_on_issues',False) and self.report['issues']:raise ContentError('EXPORT_REVIEW_REQUIRED','导出存在未解决项目',details=self.report)
        path.parent.mkdir(parents=True,exist_ok=True)
        tmp=path.with_name(path.stem+'.'+uuid.uuid4().hex+'.docx')
        try:self.doc.save(tmp);tmp.replace(path)
        finally:tmp.unlink(missing_ok=True)
        self.report['output_sha256']=hashlib.sha256(path.read_bytes()).hexdigest()
        report_path=path.with_suffix('.report.json')
        report_path.write_text(json.dumps(self.report,ensure_ascii=False,indent=2),encoding='utf-8')
        return {'output_path':str(path.resolve()),'report_path':str(report_path.resolve()),'report':self.report}
